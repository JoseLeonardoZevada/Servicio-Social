import os
import json
import re

import gradio as gr
import pdfplumber
import google.generativeai as genai
from pptx import Presentation

# La API key se lee del "Secret" que configuraste en el Space
genai.configure(api_key=os.environ["GEMINI_API_KEY"])
model = genai.GenerativeModel("gemini-2.5-flash")


def extract_text(pdf_path):
    """Saca todo el texto del PDF, página por página."""
    text_parts = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def ask_gemini_for_slides(text):
    """Le pide a Gemini que convierta el texto en una estructura de diapositivas (JSON)."""
    prompt = f"""
Eres un asistente que convierte un documento en el contenido de una presentación.
Lee el siguiente texto y devuelve EXCLUSIVAMENTE un JSON (sin texto adicional,
sin markdown, sin ``` ), con esta estructura exacta:

{{
  "slides": [
    {{"title": "Titulo de la diapositiva", "bullets": ["punto 1", "punto 2", "punto 3"]}}
  ]
}}

Genera entre 5 y 10 diapositivas que resuman lo mas importante del documento.
Cada diapositiva debe tener entre 3 y 5 puntos (bullets) cortos y claros.

Texto del documento:
\"\"\"
{text[:15000]}
\"\"\"
"""
    response = model.generate_content(prompt)
    raw = response.text.strip()

    # Por si el modelo agrega ```json ... ``` a pesar de la instrucción
    raw = re.sub(r"^```json", "", raw)
    raw = re.sub(r"^```", "", raw)
    raw = re.sub(r"```$", "", raw)
    raw = raw.strip()

    data = json.loads(raw)
    return data["slides"]


def build_pptx(slides_data, output_path):
    """Arma el archivo .pptx a partir de la estructura de diapositivas."""
    prs = Presentation()
    title_and_content_layout = prs.slide_layouts[1]

    for slide_info in slides_data:
        slide = prs.slides.add_slide(title_and_content_layout)
        slide.shapes.title.text = slide_info.get("title", "Sin titulo")

        body_placeholder = slide.placeholders[1]
        tf = body_placeholder.text_frame
        tf.clear()

        bullets = slide_info.get("bullets", [])
        for i, bullet in enumerate(bullets):
            if i == 0:
                tf.text = bullet
            else:
                p = tf.add_paragraph()
                p.text = bullet

    prs.save(output_path)


def process_pdf(pdf_file):
    if pdf_file is None:
        raise gr.Error("Por favor sube un archivo PDF.")

    text = extract_text(pdf_file)
    if not text.strip():
        raise gr.Error(
            "No se pudo extraer texto de ese PDF (¿está escaneado como imagen?)."
        )

    slides_data = ask_gemini_for_slides(text)

    output_path = "/tmp/presentacion_generada.pptx"
    build_pptx(slides_data, output_path)

    return output_path


demo = gr.Interface(
    fn=process_pdf,
    inputs=gr.File(label="Sube tu PDF", file_types=[".pdf"]),
    outputs=gr.File(label="Descarga tu presentación (.pptx)"),
    title="Agente de IA: PDF a Diapositivas",
    description="Sube un PDF y el agente generará automáticamente una presentación resumida.",
)

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 10000))
    demo.launch(server_name="0.0.0.0", server_port=port)
